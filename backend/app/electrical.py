"""Parse and generate Pablo-style BT / MV electrical configuration workbooks."""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any, BinaryIO, Iterable, Optional

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from app.models import (
    ElectricalBtConfig,
    ElectricalBtItsSummary,
    ElectricalBtRow,
    ElectricalBtTotals,
    ElectricalMvConfig,
    ElectricalMvRow,
    ProjectParameters,
)

BT_SHEET = "Electrical Configuration"
MV_SHEET = "MV Configuration"
README_SHEET = "README"

DEFAULT_MODULES_PER_TRACKER = 87
DEFAULT_STRINGS_PER_TRACKER = 3

ACCENT_FILL = PatternFill("solid", fgColor="3DCC8C")
HEADER_FONT = Font(name="IBM Plex Sans", bold=True, color="0C0E12")
BODY_FONT = Font(name="IBM Plex Sans")
MONO_FONT = Font(name="IBM Plex Mono", size=10)
README_TITLE = Font(name="IBM Plex Sans", bold=True, size=14, color="0C0E12")

BT_HEADERS = [
    "Subpark",
    "ITS",
    "ITS id",
    "Inv",
    "Inverter id",
    "Module Manufac.",
    "Module Power [W]",
    "Trackers 1Vx87",
    "Strings (Inv)",
    "Strings (ITS)",
    "Nº Modules",
    "Pot DC [kWp] (Inv)",
    "Pot DC [kWp] (ITS)",
    "Pot AC [kVA] (Inv)",
    "Pot AC [kVA] (ITS)",
    "Ratio",
    "SCB",
]

MV_HEADERS = [
    "Subpark",
    "ITS",
    "ITS id",
    "Feeder id",
    "Destination",
    "Voltage_kV",
    "Transformer_kVA",
    "Cable section_mm2",
    "Conductor",
    "Length_m",
    "Topology (radial|ring|open-ring)",
    "Load_kVA",
    "Notes",
]

_TRACKER_HEADER_RE = re.compile(
    r"(?i)trackers?\s*(?:(?P<portrait>\d+)\s*[vV]\s*[xX]\s*)?(?P<modules>\d+)?"
)
_VX_MODULES_RE = re.compile(r"(?i)[vV]\s*[xX]\s*(?P<modules>\d+)")
_FORMULA_RE = re.compile(r"^\s*=")

_BT_ALIASES: dict[str, tuple[str, ...]] = {
    "subpark": ("subpark",),
    "its": ("its",),
    "its_id": ("its id", "its_id", "itsid"),
    "inv": ("inv", "inverter", "inv no", "inv n"),
    "inverter_id": ("inverter id", "inverter_id", "inverterid"),
    "module_manufacturer": (
        "module manufac.",
        "module manufac",
        "module manufacturer",
        "manufacturer",
    ),
    "module_power_w": (
        "module power [w]",
        "module power w",
        "module power",
        "pmp w",
        "pmp [w]",
    ),
    "trackers": ("trackers", "n trackers", "nº trackers", "no trackers"),
    "modules_per_tracker_col": (
        "modules per tracker",
        "modules/tracker",
        "mod per tracker",
    ),
    "strings_per_tracker_col": (
        "strings per tracker",
        "strings/tracker",
    ),
    "strings_inv": ("strings (inv)", "strings inv", "strings inverter"),
    "strings_its": ("strings (its)", "strings its"),
    "n_modules": (
        "nº modules",
        "n modules",
        "no modules",
        "num modules",
        "modules",
    ),
    "pot_dc_kwp_inv": (
        "pot dc [kwp] (inv)",
        "pot dc kwp (inv)",
        "pdc kwp inv",
        "dc kwp inv",
    ),
    "pot_dc_kwp_its": (
        "pot dc [kwp] (its)",
        "pot dc kwp (its)",
        "pdc kwp its",
        "dc kwp its",
    ),
    "pot_ac_kva_inv": (
        "pot ac [kva] (inv)",
        "pot ac kva (inv)",
        "pac kva inv",
        "ac kva inv",
    ),
    "pot_ac_kva_its": (
        "pot ac [kva] (its)",
        "pot ac kva (its)",
        "pac kva its",
        "ac kva its",
    ),
    "ratio": ("ratio", "dc/ac", "dc ac"),
    "scb": ("scb", "string box", "string boxes"),
}

_MV_ALIASES: dict[str, tuple[str, ...]] = {
    "subpark": ("subpark",),
    "its": ("its",),
    "its_id": ("its id", "its_id", "itsid"),
    "feeder_id": ("feeder id", "feeder_id", "feeder"),
    "destination": ("destination", "dest", "to"),
    "voltage_kv": ("voltage_kv", "voltage kv", "voltage [kv]", "kv"),
    "transformer_kva": (
        "transformer_kva",
        "transformer kva",
        "transformer [kva]",
        "trx kva",
    ),
    "cable_section_mm2": (
        "cable section_mm2",
        "cable section mm2",
        "section_mm2",
        "section mm2",
    ),
    "conductor": ("conductor",),
    "length_m": ("length_m", "length m", "length [m]", "length"),
    "topology": (
        "topology (radial|ring|open-ring)",
        "topology",
    ),
    "load_kva": ("load_kva", "load kva", "load [kva]"),
    "notes": ("notes", "note", "comments"),
}


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _norm_header(value: Any) -> str:
    text = "" if value is None else str(value).strip().lower()
    text = text.replace("º", "o").replace("°", "o")
    text = re.sub(r"\s+", " ", text)
    return text


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and bool(_FORMULA_RE.match(value))


def _as_str(value: Any) -> str:
    if value is None or _is_formula(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _as_number(value: Any) -> Optional[float]:
    if value is None or value == "" or _is_formula(value):
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _coalesce(*values: Any) -> Any:
    for value in values:
        if value is None or value == "":
            continue
        if _is_formula(value):
            continue
        return value
    return None


def _sum_optional(values: Iterable[Optional[float]]) -> Optional[float]:
    total = 0.0
    seen = False
    for value in values:
        if value is None:
            continue
        total += value
        seen = True
    return total if seen else None


def _safe_div(num: Optional[float], den: Optional[float]) -> Optional[float]:
    if num is None or den in (None, 0):
        return None
    return num / den


def detect_tracker_geometry(
    header: str,
    *,
    project_modules: Optional[int] = None,
    project_strings: Optional[int] = None,
) -> tuple[int, int]:
    """Modules/strings per tracker from Pablo's Trackers 1Vx87 header or project params."""
    modules = project_modules if project_modules and project_modules > 0 else None
    strings = project_strings if project_strings and project_strings > 0 else None
    match = _TRACKER_HEADER_RE.search(header or "")
    vx = _VX_MODULES_RE.search(header or "")
    if vx and modules is None:
        modules = int(vx.group("modules"))
    elif match and match.group("modules") and modules is None:
        modules = int(match.group("modules"))
    if strings is None:
        strings = DEFAULT_STRINGS_PER_TRACKER
    if modules is None:
        modules = DEFAULT_MODULES_PER_TRACKER
    return modules, strings


def _map_headers(headers: list[Any], aliases: dict[str, tuple[str, ...]]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    normalized = [_norm_header(h) for h in headers]
    for field, names in aliases.items():
        for idx, header in enumerate(normalized):
            if not header:
                continue
            if header in names:
                mapping[field] = idx
                break
            if field == "trackers" and header.startswith("trackers"):
                mapping[field] = idx
                break
            if field == "topology" and header.startswith("topology"):
                mapping[field] = idx
                break
    return mapping


def _cell_pair(value_row: tuple[Any, ...], formula_row: tuple[Any, ...], idx: Optional[int]) -> Any:
    if idx is None or idx >= len(value_row):
        return None
    cached = value_row[idx] if idx < len(value_row) else None
    formula = formula_row[idx] if idx < len(formula_row) else None
    return _coalesce(cached, None if _is_formula(formula) else formula)


def _row_empty(values: tuple[Any, ...]) -> bool:
    return all(v is None or v == "" for v in values)


def _its_id(subpark: str, its: str, existing: str = "") -> str:
    if existing:
        return existing
    if subpark and its:
        return f"ITS {subpark}-{its}"
    return existing


def _inverter_id(subpark: str, its: str, inv: str, existing: str = "") -> str:
    if existing:
        return existing
    if subpark and its and inv:
        return f"{subpark}-{its}-{inv}"
    return existing


def _summarize_bt(rows: list[ElectricalBtRow]) -> tuple[list[ElectricalBtItsSummary], ElectricalBtTotals]:
    groups: dict[str, list[ElectricalBtRow]] = {}
    order: list[str] = []
    for row in rows:
        key = row.its_id or f"{row.subpark}|{row.its}"
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(row)

    summaries: list[ElectricalBtItsSummary] = []
    for key in order:
        group = groups[key]
        first = group[0]
        dc = _sum_optional(r.pot_dc_kwp_inv for r in group)
        ac = _sum_optional(r.pot_ac_kva_inv for r in group)
        summaries.append(
            ElectricalBtItsSummary(
                subpark=first.subpark,
                its=first.its,
                its_id=first.its_id or key,
                inverter_count=len(group),
                trackers=_sum_optional(r.trackers for r in group),
                strings=_sum_optional(r.strings_inv for r in group),
                n_modules=_sum_optional(r.n_modules for r in group),
                pot_dc_kwp=dc,
                pot_ac_kva=ac,
                ratio=_safe_div(dc, ac),
            )
        )

    totals = ElectricalBtTotals(
        its_count=len(summaries),
        inverter_count=len(rows),
        trackers=_sum_optional(r.trackers for r in rows),
        n_modules=_sum_optional(r.n_modules for r in rows),
        pot_dc_kwp=_sum_optional(r.pot_dc_kwp_inv for r in rows),
        pot_ac_kva=_sum_optional(r.pot_ac_kva_inv for r in rows),
    )
    return summaries, totals


def _apply_its_aggregates(rows: list[ElectricalBtRow], summaries: list[ElectricalBtItsSummary]) -> None:
    by_id = {s.its_id: s for s in summaries}
    for row in rows:
        summary = by_id.get(row.its_id)
        if summary is None:
            continue
        if row.strings_its is None:
            row.strings_its = summary.strings
        if row.pot_dc_kwp_its is None:
            row.pot_dc_kwp_its = summary.pot_dc_kwp
        if row.pot_ac_kva_its is None:
            row.pot_ac_kva_its = summary.pot_ac_kva


def _open_workbooks(raw: bytes):
    try:
        wb_values = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        wb_formulas = load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
    except (InvalidFileException, OSError, KeyError) as exc:
        raise ValueError(f"Could not read workbook: {exc}") from exc
    return wb_values, wb_formulas


def parse_bt_workbook(
    data: bytes | BinaryIO,
    *,
    filename: str = "",
    parameters: Optional[ProjectParameters] = None,
) -> ElectricalBtConfig:
    warnings: list[str] = []
    params = parameters or ProjectParameters()
    raw = data if isinstance(data, (bytes, bytearray)) else data.read()
    wb_values, wb_formulas = _open_workbooks(raw)

    sheet_name = BT_SHEET if BT_SHEET in wb_values.sheetnames else wb_values.sheetnames[0]
    if sheet_name != BT_SHEET:
        warnings.append(f"Sheet '{BT_SHEET}' not found; using '{sheet_name}'.")
    ws_values = wb_values[sheet_name]
    ws_formulas = wb_formulas[sheet_name] if sheet_name in wb_formulas.sheetnames else wb_formulas[wb_formulas.sheetnames[0]]

    value_rows = list(ws_values.iter_rows(values_only=True))
    formula_rows = list(ws_formulas.iter_rows(values_only=True))
    wb_values.close()
    wb_formulas.close()

    if not value_rows:
        raise ValueError("Workbook has no rows.")

    headers = list(value_rows[0])
    mapping = _map_headers(headers, _BT_ALIASES)
    if "subpark" not in mapping or "its" not in mapping:
        raise ValueError("BT sheet must include Subpark and ITS columns.")

    trackers_header = ""
    if "trackers" in mapping and mapping["trackers"] < len(headers):
        trackers_header = "" if headers[mapping["trackers"]] is None else str(headers[mapping["trackers"]])

    modules_per_tracker, strings_per_tracker = detect_tracker_geometry(
        trackers_header,
        project_modules=params.modules_per_tracker,
        project_strings=params.strings_per_tracker,
    )

    rows: list[ElectricalBtRow] = []
    for ridx in range(1, max(len(value_rows), len(formula_rows))):
        value_row = value_rows[ridx] if ridx < len(value_rows) else tuple()
        formula_row = formula_rows[ridx] if ridx < len(formula_rows) else tuple()
        if _row_empty(value_row) and _row_empty(formula_row):
            continue

        def col(field: str) -> Any:
            return _cell_pair(value_row, formula_row, mapping.get(field))

        subpark = _as_str(col("subpark"))
        its = _as_str(col("its"))
        inv = _as_str(col("inv"))
        if not subpark and not its and not inv:
            continue

        its_id = _its_id(subpark, its, _as_str(col("its_id")))
        inverter_id = _inverter_id(subpark, its, inv, _as_str(col("inverter_id")))

        mpt_override = _as_number(col("modules_per_tracker_col"))
        spt_override = _as_number(col("strings_per_tracker_col"))
        row_modules = int(mpt_override) if mpt_override and mpt_override > 0 else modules_per_tracker
        row_strings = int(spt_override) if spt_override and spt_override > 0 else strings_per_tracker

        trackers = _as_number(col("trackers"))
        module_power_w = _as_number(col("module_power_w"))
        strings_inv = _as_number(col("strings_inv"))
        n_modules = _as_number(col("n_modules"))
        pot_dc_inv = _as_number(col("pot_dc_kwp_inv"))
        pot_ac_inv = _as_number(col("pot_ac_kva_inv"))
        ratio = _as_number(col("ratio"))

        if strings_inv is None and trackers is not None:
            strings_inv = trackers * row_strings
        if n_modules is None and trackers is not None:
            n_modules = trackers * row_modules
        if pot_dc_inv is None and module_power_w is not None and n_modules is not None:
            pot_dc_inv = module_power_w * n_modules / 1000.0
        if ratio is None:
            ratio = _safe_div(pot_dc_inv, pot_ac_inv)

        if not subpark or not its:
            warnings.append(f"Row {ridx + 1}: missing Subpark or ITS key.")
        if not inv:
            warnings.append(f"Row {ridx + 1}: missing Inv.")

        scb_raw = col("scb")
        rows.append(
            ElectricalBtRow(
                subpark=subpark,
                its=its,
                its_id=its_id,
                inv=inv,
                inverter_id=inverter_id,
                module_manufacturer=_as_str(col("module_manufacturer")),
                module_power_w=module_power_w,
                trackers=trackers,
                strings_inv=strings_inv,
                strings_its=_as_number(col("strings_its")),
                n_modules=n_modules,
                pot_dc_kwp_inv=pot_dc_inv,
                pot_dc_kwp_its=_as_number(col("pot_dc_kwp_its")),
                pot_ac_kva_inv=pot_ac_inv,
                pot_ac_kva_its=_as_number(col("pot_ac_kva_its")),
                ratio=ratio,
                scb=_as_str(scb_raw) if scb_raw is not None else "",
            )
        )

    if not rows:
        raise ValueError("No inverter rows found in BT sheet.")

    by_its, totals = _summarize_bt(rows)
    _apply_its_aggregates(rows, by_its)

    return ElectricalBtConfig(
        source_filename=filename,
        imported_at=_now(),
        sheet_name=sheet_name,
        trackers_header=trackers_header,
        modules_per_tracker=modules_per_tracker,
        strings_per_tracker=strings_per_tracker,
        rows=rows,
        by_its=by_its,
        totals=totals,
        warnings=warnings,
    )


def _normalize_topology(value: str) -> Optional[str]:
    text = value.strip().lower().replace("_", "-").replace(" ", "-")
    aliases = {
        "radial": "radial",
        "ring": "ring",
        "open-ring": "open-ring",
        "openring": "open-ring",
        "open": "open-ring",
    }
    return aliases.get(text)


def parse_mv_workbook(
    data: bytes | BinaryIO,
    *,
    filename: str = "",
    bt: Optional[ElectricalBtConfig] = None,
) -> ElectricalMvConfig:
    warnings: list[str] = []
    raw = data if isinstance(data, (bytes, bytearray)) else data.read()
    wb_values, wb_formulas = _open_workbooks(raw)

    sheet_name = MV_SHEET if MV_SHEET in wb_values.sheetnames else None
    if sheet_name is None:
        for name in wb_values.sheetnames:
            if name.strip().lower() not in {"readme", "example"}:
                sheet_name = name
                break
    if sheet_name is None:
        raise ValueError("Workbook has no data sheet.")
    if sheet_name != MV_SHEET:
        warnings.append(f"Sheet '{MV_SHEET}' not found; using '{sheet_name}'.")

    ws_values = wb_values[sheet_name]
    ws_formulas = wb_formulas[sheet_name] if sheet_name in wb_formulas.sheetnames else wb_formulas[wb_formulas.sheetnames[0]]
    value_rows = list(ws_values.iter_rows(values_only=True))
    formula_rows = list(ws_formulas.iter_rows(values_only=True))
    wb_values.close()
    wb_formulas.close()

    if not value_rows:
        raise ValueError("MV sheet has no rows.")

    mapping = _map_headers(list(value_rows[0]), _MV_ALIASES)
    if "subpark" not in mapping or "its" not in mapping:
        raise ValueError("MV sheet must include Subpark and ITS columns.")

    known_its: set[str] = set()
    known_keys: set[tuple[str, str]] = set()
    if bt:
        for row in bt.rows:
            known_its.add(row.its_id)
            known_keys.add((row.subpark, row.its))

    rows: list[ElectricalMvRow] = []
    for ridx in range(1, max(len(value_rows), len(formula_rows))):
        value_row = value_rows[ridx] if ridx < len(value_rows) else tuple()
        formula_row = formula_rows[ridx] if ridx < len(formula_rows) else tuple()
        if _row_empty(value_row) and _row_empty(formula_row):
            continue

        def col(field: str) -> Any:
            return _cell_pair(value_row, formula_row, mapping.get(field))

        subpark = _as_str(col("subpark"))
        its = _as_str(col("its"))
        if not subpark and not its and not _as_str(col("feeder_id")):
            continue
        its_id = _its_id(subpark, its, _as_str(col("its_id")))
        topology_raw = _as_str(col("topology"))
        topology = _normalize_topology(topology_raw) if topology_raw else None
        if topology_raw and topology is None:
            warnings.append(f"Row {ridx + 1}: unknown topology '{topology_raw}'.")
        if not subpark or not its:
            warnings.append(f"Row {ridx + 1}: missing Subpark or ITS key.")
        if bt and its_id not in known_its and (subpark, its) not in known_keys:
            warnings.append(
                f"Row {ridx + 1}: ITS '{its_id or subpark + '-' + its}' is not in the BT configuration."
            )

        rows.append(
            ElectricalMvRow(
                subpark=subpark,
                its=its,
                its_id=its_id,
                feeder_id=_as_str(col("feeder_id")),
                destination=_as_str(col("destination")),
                voltage_kv=_as_number(col("voltage_kv")),
                transformer_kva=_as_number(col("transformer_kva")),
                cable_section_mm2=_as_number(col("cable_section_mm2")),
                conductor=_as_str(col("conductor")),
                length_m=_as_number(col("length_m")),
                topology=topology,  # type: ignore[arg-type]
                load_kva=_as_number(col("load_kva")),
                notes=_as_str(col("notes")),
            )
        )

    if not rows:
        raise ValueError("No MV feeder rows found.")

    return ElectricalMvConfig(
        source_filename=filename,
        imported_at=_now(),
        sheet_name=sheet_name,
        rows=rows,
        warnings=warnings,
    )


def merge_bt(existing: Optional[ElectricalBtConfig], incoming: ElectricalBtConfig) -> ElectricalBtConfig:
    if existing is None or not existing.rows:
        return incoming
    by_key: dict[str, ElectricalBtRow] = {}
    order: list[str] = []

    def key_of(row: ElectricalBtRow) -> str:
        return row.inverter_id or f"{row.subpark}|{row.its}|{row.inv}"

    for row in existing.rows:
        k = key_of(row)
        by_key[k] = row
        order.append(k)
    for row in incoming.rows:
        k = key_of(row)
        if k not in by_key:
            order.append(k)
        by_key[k] = row
    rows = [by_key[k] for k in order]
    by_its, totals = _summarize_bt(rows)
    _apply_its_aggregates(rows, by_its)
    return incoming.model_copy(
        update={
            "rows": rows,
            "by_its": by_its,
            "totals": totals,
            "warnings": list(dict.fromkeys([*existing.warnings, *incoming.warnings])),
        }
    )


def merge_mv(existing: Optional[ElectricalMvConfig], incoming: ElectricalMvConfig) -> ElectricalMvConfig:
    if existing is None or not existing.rows:
        return incoming

    def key_of(row: ElectricalMvRow) -> str:
        return "|".join([row.its_id or f"{row.subpark}-{row.its}", row.feeder_id, row.destination])

    by_key: dict[str, ElectricalMvRow] = {}
    order: list[str] = []
    for row in existing.rows:
        k = key_of(row)
        by_key[k] = row
        order.append(k)
    for row in incoming.rows:
        k = key_of(row)
        if k not in by_key:
            order.append(k)
        by_key[k] = row
    return incoming.model_copy(
        update={
            "rows": [by_key[k] for k in order],
            "warnings": list(dict.fromkeys([*existing.warnings, *incoming.warnings])),
        }
    )


def _style_header(ws: Worksheet, headers: list[str], widths: list[int]) -> None:
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(1, idx, header)
        cell.fill = ACCENT_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(idx)].width = widths[idx - 1] if idx - 1 < len(widths) else 14
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _write_readme(ws: Worksheet, lines: list[str]) -> None:
    ws.column_dimensions["A"].width = 110
    ws["A1"] = lines[0]
    ws["A1"].font = README_TITLE
    for idx, line in enumerate(lines[1:], start=2):
        ws.cell(idx, 1, line).font = BODY_FONT
        ws.cell(idx, 1).alignment = Alignment(wrap_text=True)


def build_bt_template(*, include_examples: bool = True) -> bytes:
    wb = Workbook()
    readme = wb.active
    readme.title = README_SHEET
    _write_readme(
        readme,
        [
            "BT / LV Electrical Configuration",
            "Fill the Electrical Configuration sheet. One row = one inverter (or inverter slice) under an ITS.",
            "Several rows may share the same ITS. Join key for MV is Subpark + ITS (and ITS id).",
            "Typical inputs: Subpark, ITS, Inv, Module Manufac., Module Power [W], Trackers count, Pot AC [kVA] (Inv), optional SCB.",
            "ITS id, Inverter id, strings, modules, DC power, ITS aggregates, and Ratio are formulas in this template.",
            "Trackers header 'Trackers 1Vx87' means 87 modules per tracker and 3 strings per tracker (H*3 strings, H*87 modules).",
            "If you use a different table (e.g. 1Vx56), rename that column or add a 'Modules per tracker' column.",
            "Importer stores cached values when present, and recomputes the obvious formulas when only inputs are filled.",
        ],
    )

    ws = wb.create_sheet(BT_SHEET)
    headers = list(BT_HEADERS)
    _style_header(ws, headers, [12, 10, 14, 10, 14, 18, 18, 16, 14, 14, 14, 18, 18, 18, 18, 12, 10])
    example_count = 4 if include_examples else 0
    last = max(example_count, 1) + 1
    for r in range(2, 2 + max(example_count, 2)):
        ws.cell(r, 3, f'=+"ITS "&A{r}&"-"&B{r}').font = MONO_FONT
        ws.cell(r, 5, f'=+A{r}&"-"&B{r}&"-"&D{r}').font = MONO_FONT
        ws.cell(r, 9, f"=+H{r}*3").font = MONO_FONT
        ws.cell(r, 10, f"=SUMIF($C$2:$C${last},C{r},$I$2:$I${last})").font = MONO_FONT
        ws.cell(r, 11, f"=+H{r}*87").font = MONO_FONT
        ws.cell(r, 12, f"=+G{r}*K{r}/1000").font = MONO_FONT
        ws.cell(r, 13, f"=SUMIF($C$2:$C${last},C{r},$L$2:$L${last})").font = MONO_FONT
        ws.cell(r, 15, f"=SUMIF($C$2:$C${last},C{r},$N$2:$N${last})").font = MONO_FONT
        ws.cell(r, 16, f"=IF(N{r}=0,\"\",L{r}/N{r})").font = MONO_FONT
    if include_examples:
        examples = [
            (1, 1, 1, "CHINT", 345, 133, 3682),
            (1, 1, 2, "CHINT", 345, 163, 3682),
            (2, 2, 1, "CHINT", 345, 117, 3682),
            (2, 2, 2, "CHINT", 345, 178, 3682),
        ]
        for r, (sub, its, inv, mfr, pmp, trackers, pac) in enumerate(examples, start=2):
            ws.cell(r, 1, sub)
            ws.cell(r, 2, its)
            ws.cell(r, 4, inv)
            ws.cell(r, 6, mfr)
            ws.cell(r, 7, pmp)
            ws.cell(r, 8, trackers)
            ws.cell(r, 14, pac)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_mv_template(*, include_examples: bool = True) -> bytes:
    wb = Workbook()
    readme = wb.active
    readme.title = README_SHEET
    _write_readme(
        readme,
        [
            "MV Electrical Configuration",
            "One row = one MV feeder (or cable) leaving an ITS / PCS.",
            "Join to the BT sheet on Subpark + ITS, or on ITS id (e.g. ITS 1-1).",
            "Topology: radial | ring | open-ring.",
            "Destination is typically the MV bus, substation, or next ITS in a ring.",
            "This workbook does not replace plant layout stations; it stores the MV table on the project.",
        ],
    )
    ws = wb.create_sheet(MV_SHEET)
    _style_header(
        ws,
        MV_HEADERS,
        [12, 10, 14, 14, 18, 12, 18, 18, 14, 12, 28, 12, 28],
    )
    if include_examples:
        examples = [
            (1, 1, "ITS 1-1", "F-1-1", "MV bus", 33, 8000, 240, "Al", 850, "radial", 7364, "Example feeder"),
            (2, 2, "ITS 2-2", "F-2-2", "MV bus", 33, 8000, 240, "Al", 920, "radial", 7364, ""),
        ]
        for r, row in enumerate(examples, start=2):
            for c, value in enumerate(row, start=1):
                ws.cell(r, c, value)
            ws.cell(r, 3, f'=+"ITS "&A{r}&"-"&B{r}')
    else:
        ws.cell(2, 3, '="ITS "&A2&"-"&B2')

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def preview_bt(config: ElectricalBtConfig) -> dict[str, Any]:
    t = config.totals
    return {
        "kind": "bt",
        "its_count": t.its_count,
        "inverter_count": t.inverter_count,
        "trackers": t.trackers,
        "n_modules": t.n_modules,
        "pot_dc_kwp": t.pot_dc_kwp,
        "pot_ac_kva": t.pot_ac_kva,
        "modules_per_tracker": config.modules_per_tracker,
        "strings_per_tracker": config.strings_per_tracker,
        "warnings": config.warnings,
    }


def preview_mv(config: ElectricalMvConfig) -> dict[str, Any]:
    load = _sum_optional(r.load_kva for r in config.rows)
    length = _sum_optional(r.length_m for r in config.rows)
    its_ids = {r.its_id for r in config.rows if r.its_id}
    return {
        "kind": "mv",
        "feeder_count": len(config.rows),
        "its_count": len(its_ids),
        "load_kva": load,
        "length_m": length,
        "warnings": config.warnings,
    }
