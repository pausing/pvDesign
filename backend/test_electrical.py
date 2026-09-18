"""BT / MV electrical Excel parse, templates, and project-scoped routes."""

from __future__ import annotations

import io
import os
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

os.environ["PVDES_DATA_DIR"] = "/tmp/pvdes_electrical_test_data"
os.environ["PV_BASE_PATH"] = "/pv"

from app.electrical import (  # noqa: E402
    build_bt_template,
    build_mv_template,
    parse_bt_workbook,
    parse_mv_workbook,
)
from app.main import app  # noqa: E402
from app.models import ProjectParameters  # noqa: E402

FIXTURE = Path(__file__).resolve().parent / "tests" / "fixtures" / "electrical_bt_sample.xlsx"
IDENTITY = {
    "X-Powerlearn-User-Id": "user-42",
    "X-Powerlearn-Email": "pablo@powerlearn.us",
    "X-Powerlearn-Admin": "true",
}


@pytest.fixture
def client():
    return TestClient(app)


def _create_project(client: TestClient, headers: dict | None = None) -> str:
    response = client.post(
        "/pv/api/projects",
        json={"name": "Electrical", "site": "Lab", "seed_catalog": False},
        headers=headers or {},
    )
    assert response.status_code == 201
    return response.json()["id"]


def test_parse_pablo_sample_fixture():
    assert FIXTURE.is_file()
    config = parse_bt_workbook(FIXTURE.read_bytes(), filename=FIXTURE.name)
    assert config.totals.inverter_count == 36
    assert config.totals.its_count == 15
    assert config.modules_per_tracker == 87
    assert config.strings_per_tracker == 3
    assert config.trackers_header == "Trackers 1Vx87"
    first = config.rows[0]
    assert first.its_id == "ITS 1-1"
    assert first.inverter_id == "1-1-1"
    assert first.module_manufacturer == "CHINT"
    assert first.module_power_w == 345
    assert first.trackers == 133
    assert first.strings_inv == 399
    assert first.n_modules == 11571
    assert first.pot_dc_kwp_inv == pytest.approx(3991.995)
    assert first.pot_ac_kva_inv == 3682
    assert config.totals.pot_dc_kwp == pytest.approx(sum(r.pot_dc_kwp_inv or 0 for r in config.rows))
    its1 = next(s for s in config.by_its if s.its_id == "ITS 1-1")
    assert its1.inverter_count == 2
    assert its1.pot_ac_kva == pytest.approx(7364)


def test_parse_formula_only_inputs():
    wb = Workbook()
    ws = wb.active
    ws.title = "Electrical Configuration"
    ws.append(
        [
            "Subpark",
            "ITS",
            "ITS id",
            "Inv",
            "Inverter id",
            "Module Manufac.",
            "Module Power [W]",
            "Trackers 1Vx87",
            "Strings (Inv)",
            "Strings (ITS)",
            "Nº Modules",
            "Pot DC [kWp] (Inv)",
            "Pot DC [kWp] (ITS)",
            "Pot AC [kVA] (Inv)",
            "Pot AC [kVA] (ITS)",
            "Ratio",
            "SCB",
        ]
    )
    ws.append([1, 1, "=+\"ITS \"&A2&\"-\"&B2", 1, "=+A2&\"-\"&B2&\"-\"&D2", "CHINT", 345, 10, "=H2*3", None, "=H2*87", "=G2*K2/1000", None, 1000, None, "=L2/N2", None])
    ws.append([1, 1, "=+\"ITS \"&A3&\"-\"&B3", 2, "=+A3&\"-\"&B3&\"-\"&D3", "CHINT", 345, 20, "=H3*3", None, "=H3*87", "=G3*K3/1000", None, 1000, None, "=L3/N3", None])
    buf = io.BytesIO()
    wb.save(buf)
    config = parse_bt_workbook(buf.getvalue(), filename="formulas.xlsx")
    assert config.totals.inverter_count == 2
    assert config.rows[0].its_id == "ITS 1-1"
    assert config.rows[0].inverter_id == "1-1-1"
    assert config.rows[0].strings_inv == 30
    assert config.rows[0].n_modules == 870
    assert config.rows[0].pot_dc_kwp_inv == pytest.approx(300.15)
    assert config.rows[1].n_modules == 1740
    assert config.by_its[0].pot_dc_kwp == pytest.approx(300.15 + 600.3)
    assert config.rows[0].strings_its == 90
    assert config.rows[0].ratio == pytest.approx(0.30015)


def test_modules_per_tracker_from_project_parameter_not_hardcoded():
    wb = Workbook()
    ws = wb.active
    ws.title = "Electrical Configuration"
    ws.append(
        [
            "Subpark",
            "ITS",
            "Inv",
            "Module Manufac.",
            "Module Power [W]",
            "Trackers",
            "Pot AC [kVA] (Inv)",
            "Modules per tracker",
        ]
    )
    ws.append([1, 1, 1, "Jinko", 580, 2, 3000, 56])
    buf = io.BytesIO()
    wb.save(buf)
    config = parse_bt_workbook(
        buf.getvalue(),
        parameters=ProjectParameters(modules_per_tracker=56, strings_per_tracker=2),
    )
    assert config.modules_per_tracker == 56
    assert config.strings_per_tracker == 2
    assert config.rows[0].n_modules == 112
    assert config.rows[0].strings_inv == 4
    assert config.rows[0].its_id == "ITS 1-1"


def test_bt_and_mv_templates_have_readme_and_headers():
    bt = load_workbook(io.BytesIO(build_bt_template()), data_only=False)
    assert "README" in bt.sheetnames
    assert "Electrical Configuration" in bt.sheetnames
    headers = [c.value for c in bt["Electrical Configuration"][1]]
    assert headers[0] == "Subpark"
    assert headers[7] == "Trackers 1Vx87"
    assert bt["Electrical Configuration"]["C2"].value.startswith("=")
    assert bt["Electrical Configuration"]["A2"].value == 1

    mv = load_workbook(io.BytesIO(build_mv_template()), data_only=False)
    assert "README" in mv.sheetnames
    assert "MV Configuration" in mv.sheetnames
    mv_headers = [c.value for c in mv["MV Configuration"][1]]
    assert mv_headers[3] == "Feeder id"
    assert "Topology" in mv_headers[10]


def test_parse_mv_joins_its_and_warns_unknown(client):
    bt = parse_bt_workbook(FIXTURE.read_bytes())
    mv_bytes = build_mv_template(include_examples=True)
    config = parse_mv_workbook(mv_bytes, bt=bt)
    assert len(config.rows) == 2
    assert config.rows[0].its_id == "ITS 1-1"
    assert config.rows[0].topology == "radial"
    assert not any("not in the BT" in w for w in config.warnings)

    wb = Workbook()
    ws = wb.active
    ws.title = "MV Configuration"
    ws.append(
        [
            "Subpark",
            "ITS",
            "ITS id",
            "Feeder id",
            "Destination",
            "Voltage_kV",
            "Transformer_kVA",
            "Cable section_mm2",
            "Conductor",
            "Length_m",
            "Topology (radial|ring|open-ring)",
            "Load_kVA",
            "Notes",
        ]
    )
    ws.append([99, 99, None, "F-x", "Bus", 33, 8000, 240, "Al", 10, "ring", 100, ""])
    buf = io.BytesIO()
    wb.save(buf)
    unknown = parse_mv_workbook(buf.getvalue(), bt=bt)
    assert unknown.rows[0].its_id == "ITS 99-99"
    assert any("not in the BT" in w for w in unknown.warnings)


def test_template_download_requires_project(client):
    missing = client.get("/pv/api/projects/nope/electrical/templates/bt")
    assert missing.status_code == 404
    project_id = _create_project(client)
    ok = client.get(f"/pv/api/projects/{project_id}/electrical/templates/bt")
    assert ok.status_code == 200
    assert "spreadsheetml" in ok.headers.get("content-type", "")
    mv = client.get(f"/pv/api/projects/{project_id}/electrical/templates/mv")
    assert mv.status_code == 200
    client.delete(f"/pv/api/projects/{project_id}")


def test_project_scoped_electrical_routes_keep_portal_identity(client):
    """Identity headers still work; unauthenticated-local access is unchanged (no 401)."""
    anonymous_id = _create_project(client)
    authed_id = _create_project(client, headers=IDENTITY)

    anon_tpl = client.get(f"/pv/api/projects/{anonymous_id}/electrical/templates/bt")
    assert anon_tpl.status_code == 200

    authed_tpl = client.get(
        f"/pv/api/projects/{authed_id}/electrical/templates/mv",
        headers=IDENTITY,
    )
    assert authed_tpl.status_code == 200

    with FIXTURE.open("rb") as fh:
        imported = client.post(
            f"/pv/api/projects/{authed_id}/electrical/bt/import",
            files={"file": ("electrical_bt_sample.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=IDENTITY,
        )
    assert imported.status_code == 200
    body = imported.json()
    assert body["preview"]["its_count"] == 15
    assert body["preview"]["inverter_count"] == 36
    fetched = client.get(f"/pv/api/projects/{authed_id}", headers=IDENTITY)
    assert fetched.status_code == 200
    assert fetched.json()["electrical_bt"]["totals"]["its_count"] == 15
    assert fetched.json()["owner"] == "pablo@powerlearn.us"
    assert fetched.json()["user_id"] == "user-42"

    me = client.get("/pv/api/me", headers=IDENTITY)
    assert me.json()["email"] == "pablo@powerlearn.us"

    client.delete(f"/pv/api/projects/{anonymous_id}")
    client.delete(f"/pv/api/projects/{authed_id}")


def test_import_bt_replace_and_merge(client):
    project_id = _create_project(client)
    with FIXTURE.open("rb") as fh:
        first = client.post(
            f"/pv/api/projects/{project_id}/electrical/bt/import?mode=replace",
            files={"file": ("electrical_bt_sample.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert first.status_code == 200
    assert first.json()["preview"]["inverter_count"] == 36

    extra = Workbook()
    ws = extra.active
    ws.title = "Electrical Configuration"
    ws.append(["Subpark", "ITS", "Inv", "Module Manufac.", "Module Power [W]", "Trackers 1Vx87", "Pot AC [kVA] (Inv)"])
    ws.append([1, 1, 1, "CHINT", 345, 200, 3682])
    ws.append([90, 1, 1, "GCL", 350, 5, 1000])
    buf = io.BytesIO()
    extra.save(buf)
    merged = client.post(
        f"/pv/api/projects/{project_id}/electrical/bt/import?mode=merge",
        files={"file": ("delta.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert merged.status_code == 200
    rows = merged.json()["electrical_bt"]["rows"]
    assert any(r["inverter_id"] == "90-1-1" for r in rows)
    updated = next(r for r in rows if r["inverter_id"] == "1-1-1")
    assert updated["trackers"] == 200
    assert len(rows) == 37
    client.delete(f"/pv/api/projects/{project_id}")


def test_parse_endpoint_does_not_persist(client):
    project_id = _create_project(client)
    with FIXTURE.open("rb") as fh:
        parsed = client.post(
            f"/pv/api/projects/{project_id}/electrical/bt/parse",
            files={"file": ("electrical_bt_sample.xlsx", fh, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        )
    assert parsed.status_code == 200
    assert parsed.json()["preview"]["its_count"] == 15
    stored = client.get(f"/pv/api/projects/{project_id}").json()
    assert stored.get("electrical_bt") is None
    client.delete(f"/pv/api/projects/{project_id}")


def test_import_mv_from_template(client):
    project_id = _create_project(client)
    tpl = client.get(f"/pv/api/projects/{project_id}/electrical/templates/mv")
    imported = client.post(
        f"/pv/api/projects/{project_id}/electrical/mv/import",
        files={"file": ("mv.xlsx", tpl.content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200
    assert imported.json()["preview"]["feeder_count"] == 2
    client.delete(f"/pv/api/projects/{project_id}")
